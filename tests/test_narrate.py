import asyncio
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

import narrate
import providers


class DummyProvider(providers.TTSProvider):
    id = "dummy"
    label = "Dummy"

    def capabilities(self):
        return providers.Capabilities(provider_id=self.id)

    async def list_voices(self):
        return []

    async def synthesize(self, text, voice_id, out_path, *, rate, pitch):
        out_path.write_bytes(text.encode("utf-8"))


class ExtractionTests(unittest.TestCase):
    def test_docx_preserves_table_position_in_document_order(self):
        import docx

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "sample.docx"
            document = docx.Document()
            document.add_heading("Title", level=1)
            document.add_paragraph("Before table.")
            table = document.add_table(rows=2, cols=2)
            table.rows[0].cells[0].text = "Term"
            table.rows[0].cells[1].text = "Meaning"
            table.rows[1].cells[0].text = "API"
            table.rows[1].cells[1].text = "interface"
            document.add_paragraph("After table.")
            document.save(path)
            self.assertEqual(
                [block["text"] for block in narrate.extract_docx(path)],
                ["Title", "Before table.", "API. interface.", "After table."],
            )

    def test_pdf_collapses_soft_wrapped_text(self):
        page = SimpleNamespace(extract_text=lambda: "First line\ncontinues.\n\nSecond.")
        with mock.patch("pypdf.PdfReader", return_value=SimpleNamespace(pages=[page])):
            blocks = narrate.extract_pdf(Path("sample.pdf"))
        self.assertEqual(
            [block["text"] for block in blocks],
            ["First line continues.", "Second."],
        )

    def test_markdown_preserves_headings_paragraphs_and_table_rows(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "sample.md"
            path.write_text(
                "# Title\n\nOpening text.\n\n"
                "| Term | Meaning |\n| --- | --- |\n| API | interface |\n",
                encoding="utf-8",
            )
            self.assertEqual(
                narrate.extract_md(path),
                [
                    {"type": "heading", "level": 1, "text": "Title"},
                    {"type": "para", "level": 0, "text": "Opening text."},
                    {"type": "para", "level": 0, "text": "API. interface."},
                ],
            )

    def test_text_paragraphs_collapse_soft_wrapped_lines(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "sample.txt"
            path.write_text("First line\ncontinues.\n\nSecond paragraph.", encoding="utf-8")
            self.assertEqual(
                [block["text"] for block in narrate.extract_txt(path)],
                ["First line continues.", "Second paragraph."],
            )


class CleanupAndChunkingTests(unittest.TestCase):
    def test_cleanup_applies_rules_and_verbalizes_symbols(self):
        result = narrate.clean_for_speech(
            "**NOTE:** X = 50% → Y", [("NOTE:", "Important.")]
        )
        self.assertEqual(result, "Important. X equals 50 percent leads to Y")

    def test_major_headings_get_pause_between_sections(self):
        paragraphs = narrate.blocks_to_paragraphs(
            [
                {"type": "heading", "level": 1, "text": "One"},
                {"type": "para", "level": 0, "text": "Body"},
                {"type": "heading", "level": 1, "text": "Two"},
            ],
            [],
        )
        self.assertEqual(paragraphs, ["One.", "Body", narrate.PAUSE_CUE, "Two."])

    def test_chunking_prefers_sentence_boundaries(self):
        chunks = narrate.chunk_paragraphs(
            ["First sentence. Second sentence. Third sentence."], limit=32
        )
        self.assertEqual(chunks, ["First sentence. Second sentence.", "Third sentence."])

    def test_rule_file_and_cli_rules_preserve_order(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "rules.json"
            path.write_text(json.dumps({"A": "B"}), encoding="utf-8")
            self.assertEqual(
                narrate.parse_rules(["C=D"], str(path)),
                [("A", "B"), ("C", "D")],
            )


class CacheTests(unittest.TestCase):
    def test_fingerprint_changes_with_audio_affecting_settings(self):
        provider = DummyProvider()
        settings = {
            "provider": provider,
            "voice": "v",
            "rate": "+0%",
            "pitch": "+0Hz",
            "max_chars": 3000,
        }
        base = narrate.narration_fingerprint(["hello"], **settings)
        variants = [
            (["hello changed"], settings),
            (["hello"], {**settings, "voice": "other"}),
            (["hello"], {**settings, "rate": "-10%"}),
            (["hello"], {**settings, "pitch": "+5Hz"}),
            (["hello"], {**settings, "max_chars": 2000}),
        ]
        for chunks, variant in variants:
            changed = narrate.narration_fingerprint(chunks, **variant)
            self.assertNotEqual(base, changed)

    def test_cache_invalidates_stale_generated_files(self):
        with tempfile.TemporaryDirectory() as tmp:
            work_dir = Path(tmp)
            self.assertEqual(
                narrate.prepare_chunk_cache(work_dir, "first"), "initialized"
            )
            chunk = work_dir / "chunk_0001.mp3"
            chunk.write_bytes(b"old")
            self.assertEqual(narrate.prepare_chunk_cache(work_dir, "first"), "reused")
            self.assertTrue(chunk.exists())
            self.assertEqual(
                narrate.prepare_chunk_cache(work_dir, "second"), "invalidated"
            )
            self.assertFalse(chunk.exists())

    def test_generate_chunks_resumes_nonempty_files(self):
        async def run():
            with tempfile.TemporaryDirectory() as tmp:
                work_dir = Path(tmp)
                existing = work_dir / "chunk_0001.mp3"
                existing.write_bytes(b"existing")
                events = []
                paths = await narrate.generate_chunks(
                    ["one", "two"], "v", "+0%", "+0Hz", work_dir,
                    on_chunk=lambda *args: events.append(args),
                    provider=DummyProvider(),
                )
                self.assertEqual(existing.read_bytes(), b"existing")
                self.assertEqual(paths[1].read_bytes(), b"two")
                self.assertEqual(events[0][-1], "skipped")
                self.assertEqual(events[1][-1], "done")

        asyncio.run(run())


class ProviderConversionTests(unittest.TestCase):
    def test_rate_conversions_are_bounded_and_directional(self):
        self.assertLess(providers._rate_str_to_speed("-10%"), 1.0)
        self.assertGreater(providers._rate_str_to_speed("+10%"), 1.0)
        self.assertGreater(providers._rate_str_to_length_scale("-10%"), 1.0)
        self.assertLess(providers._rate_str_to_length_scale("+10%"), 1.0)
        self.assertEqual(providers._rate_str_to_wpm("+0%"), 180)
        self.assertEqual(providers._rate_str_to_speed_kokoro("+500%"), 2.0)


class AsyncPipelineTests(unittest.IsolatedAsyncioTestCase):
    async def test_async_pipeline_runs_without_nested_event_loop(self):
        provider = DummyProvider()

        def fake_encode(chunk_paths, output, work_dir, fmt):
            output.write_bytes(b"".join(path.read_bytes() for path in chunk_paths))

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.txt"
            output = root / "book.mp3"
            source.write_text("Hello.", encoding="utf-8")
            with (
                mock.patch.object(providers, "get_provider", return_value=provider),
                mock.patch.object(narrate, "encode_output", side_effect=fake_encode),
                mock.patch.object(narrate, "probe_duration", return_value=1.0),
            ):
                result = await narrate.build_audiobook_async(
                    source, output, provider="dummy", keep_chunks=True
                )
            self.assertEqual(result.output.read_bytes(), b"Hello.")
            self.assertEqual(result.duration, 1.0)


if __name__ == "__main__":
    unittest.main()
